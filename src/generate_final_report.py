#!/usr/bin/env python3
"""
Generate Final Report for WeatherCast — ML Case Study 75: Weather Pattern Analysis
Creates:
  1. docs/WeatherCast_Final_Report.docx (Word Document with figures and tables)
  2. docs/WeatherCast_Final_Report.html (Print-ready A4 HTML report with CSS and Print-to-PDF button)
"""

import os, base64
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DOCS_DIR = PROJECT_ROOT / "docs"
FIG_DIR = PROJECT_ROOT / "reports" / "figures"
DOCS_DIR.mkdir(parents=True, exist_ok=True)

# ─────────────────────────────────────────────────────────────────────────────
# 1. DOCX GENERATION
# ─────────────────────────────────────────────────────────────────────────────

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def create_docx_report():
    doc = docx.Document()
    
    # Page Setup (A4, 1-inch margins)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        
        # Header / Footer
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("WeatherCast — ML Case Study 75 | Page ")
        f_run.font.size = Pt(9)
        f_run.font.color.rgb = RGBColor(128, 128, 128)

    # Style defaults
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(51, 51, 51)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # Title Page / Header
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(4)
    run_t = title_p.add_run("WeatherCast 🌦️")
    run_t.font.name = 'Calibri'
    run_t.font.size = Pt(24)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(21, 101, 192)  # Navy Blue

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("Maharashtra Weather Pattern Analysis & Next-Day Rain Prediction\nML Case Study 75 — Final Project Report")
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(80, 80, 80)

    # Meta Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Course / Academic Level", "B.Tech Semester V — Machine Learning"),
        ("Project Identifier", "Case Study 75 (Weather Pattern Analysis)"),
        ("Target Variable", "rain_tomorrow (Binary: 1 = Rain, 0 = No Rain)"),
        ("Geographic Scope", "8 Representative Locations across Maharashtra, India"),
    ]
    for i, (k, v) in enumerate(meta_data):
        cell_k, cell_v = meta_table.rows[i].cells
        cell_k.text = k
        cell_v.text = v
        cell_k.paragraphs[0].runs[0].font.bold = True
        cell_k.paragraphs[0].runs[0].font.size = Pt(10)
        cell_v.paragraphs[0].runs[0].font.size = Pt(10)
        set_cell_background(cell_k, "F0F4F8")
        set_cell_background(cell_v, "FFFFFF")
        set_cell_margins(cell_k, 60, 60, 100, 100)
        set_cell_margins(cell_v, 60, 60, 100, 100)
    
    doc.add_paragraph() # Spacer

    # Section Helper
    def add_section_header(num, title):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(f"{num}. {title}")
        run.font.name = 'Calibri'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(21, 101, 192)
        return h

    # 1. ABSTRACT
    add_section_header("1", "Abstract")
    p = doc.add_paragraph(
        "Weather forecasting plays a vital role in agriculture, water management, disaster preparedness, and "
        "daily urban planning across Maharashtra. This case study develops WeatherCast, an end-to-end machine learning "
        "pipeline for next-day rain prediction (binary classification: rain_tomorrow = 1/0) across 8 representative "
        "locations in Maharashtra over a 25-year historical period (2000–2024, 72,691 observations). "
        "The data was collected from the Open-Meteo Historical Weather API based on the ECMWF ERA5 reanalysis model. "
        "A zero-leakage feature-engineering pipeline transformed raw atmospheric variables into 35 base features, "
        "including cyclical temporal encodings, multi-day lag metrics, and strictly backward-looking rolling statistics. "
        "Three models were developed and tuned using chronological splits (Train: 2000–2022, Validation: 2023, Test: 2024): "
        "Logistic Regression, Random Forest, and Gradient Boosting. Gradient Boosting achieved the highest overall performance "
        "with a Validation ROC-AUC of 0.9499, Test ROC-AUC of 0.9714, Test F1-score of 0.9014, and superior probability calibration "
        "(Test Brier score: 0.0665). An interactive Streamlit web application was developed where users select a city, "
        "automatically fetching live atmospheric observations to generate real-time rain predictions and calibrated confidence scores."
    )

    # 2. INTRODUCTION
    add_section_header("2", "Introduction")
    doc.add_paragraph(
        "Precipitation prediction is one of the most challenging problems in atmospheric science due to highly dynamic "
        "thermodynamic interactions, terrain-induced (orographic) effects, and localized convective instability. "
        "In Maharashtra, accurate next-day rainfall forecasting is critical because rainfall exhibits drastic spatial "
        "variability—ranging from over 3,000 mm annually in the coastal Konkan belt to under 600 mm in the central rain-shadow "
        "plateau of Madhya Maharashtra and Marathwada."
    )
    doc.add_paragraph(
        "While traditional numerical weather prediction (NWP) models solve complex fluid dynamics equations at high computational "
        "expense, machine learning offers an efficient, data-driven alternative by learning predictive atmospheric precursors "
        "(such as sudden pressure drops, humidity spikes, and accumulated multi-day surface wetness). "
        "The primary objective of this case study is to build a rigorous, leak-free supervised classification workflow "
        "that accurately predicts next-day rainfall for any selected Maharashtra city and deploys it as a user-friendly application."
    )

    # 3. PROBLEM DEFINITION
    add_section_header("3", "Problem Definition")
    doc.add_paragraph(
        "The machine learning task is formulated as a Supervised Binary Classification problem:\n"
        "• Input (X): Historical and current daily atmospheric observations (temperature, humidity, surface pressure, "
        "wind speed, precipitation hours, sunshine duration, weather codes, coordinates, cyclical date projections, and lag/rolling statistics).\n"
        "• Output (y): Target variable rain_tomorrow ∈ {0, 1}, where:\n"
        "    - 1 (Rain): Measurable rainfall occurs on the next calendar day (rain_sum > 0.0 mm).\n"
        "    - 0 (No Rain): Zero precipitation occurs on the next calendar day (rain_sum == 0.0 mm)."
    )

    # 4. DATASET & DATA SOURCE
    add_section_header("4", "Dataset & Data Source")
    doc.add_paragraph(
        "The dataset was collected from the Open-Meteo Historical Weather API using the ECMWF ERA5 reanalysis model. "
        "ERA5 provides high-resolution (0.25° grid), globally consistent historical climate data spanning 25 complete years "
        "(2000-01-01 to 2024-12-31) across 8 geographically and climatically diverse cities in Maharashtra."
    )
    
    # Dataset Table
    ds_table = doc.add_table(rows=9, cols=4)
    ds_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    ds_headers = ["Location", "Region", "Latitude / Longitude", "Rainy Days % (2000–2024)"]
    for j, h in enumerate(ds_headers):
        cell = ds_table.rows[0].cells[j]
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(9.5)
        set_cell_background(cell, "1565C0")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell, 80, 80, 100, 100)

    city_rows = [
        ("Mumbai", "Konkan (Coastal)", "19.07°N, 72.88°E", "43.4%"),
        ("Ratnagiri", "Konkan (Coastal)", "16.99°N, 73.31°E", "46.8%"),
        ("Pune", "Madhya Maharashtra", "18.52°N, 73.86°E", "44.9%"),
        ("Nashik", "Madhya Maharashtra", "20.00°N, 73.79°E", "42.3%"),
        ("Kolhapur", "Madhya Maharashtra", "16.70°N, 74.23°E", "45.1%"),
        ("Solapur", "Madhya Maharashtra", "17.67°N, 75.91°E", "35.8%"),
        ("Chhatrapati Sambhajinagar", "Marathwada", "19.88°N, 75.34°E", "37.5%"),
        ("Nagpur", "Vidarbha", "21.15°N, 79.08°E", "41.9%"),
    ]
    for i, row in enumerate(city_rows, start=1):
        for j, val in enumerate(row):
            cell = ds_table.rows[i].cells[j]
            cell.text = val
            cell.paragraphs[0].runs[0].font.size = Pt(9)
            set_cell_background(cell, "F9FBFD" if i % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, 60, 60, 100, 100)

    doc.add_paragraph(
        "Data Continuity & Target Integrity: A strict calendar verification rule was enforced: target rain_tomorrow "
        "is formed only when next_date == current_date + 1 day. In Chhatrapati Sambhajinagar, the historical record had a "
        "1-year gap during 2002 (from 2001-12-31 to 2003-01-01). The target on 2001-12-31 was explicitly assigned NaN to "
        "prevent predicting across a 366-day gap. Similarly, 2024-12-31 targets were assigned NaN across all locations. "
        "The overall raw dataset contains 72,691 rows with zero duplicate records."
    )

    # 5. DATA PREPROCESSING & FEATURE ENGINEERING
    add_section_header("5", "Data Preprocessing & Feature Engineering")
    doc.add_paragraph(
        "To ensure zero data leakage and optimal algorithmic learning, the raw observations were transformed into a supervised "
        "ML feature matrix through a calendar-strict preprocessing pipeline:"
    )
    doc.add_paragraph(
        "1. Chronological Splitting: Data was partitioned by time to mirror real-world deployment:\n"
        "   • Training Partition: 2000–2022 (66,779 observations, ~92%)\n"
        "   • Validation Partition: 2023 (2,920 observations, ~4%)\n"
        "   • Final Unseen Test Partition: 2024 (2,920 observations, ~4%)\n"
        "2. Cyclical Temporal Projections: Month and Day-of-Year integers were transformed into continuous sine and cosine "
        "projections (month_sin, month_cos, day_of_year_sin, day_of_year_cos) to eliminate artificial numerical boundaries between December and January.\n"
        "3. Calendar-Strict Lags: 8 lag features (t-1, t-3, t-7 for rain; t-1 for max/min temperature, mean humidity, pressure, and wind speed) "
        "were computed by reindexing to a complete daily calendar grid, guaranteeing that missing dates never bridged across gaps.\n"
        "4. Backward Trailing Rolling Windows: 6 rolling aggregations (3-day total rain, 7-day total rain, 3-day mean temp max/min, "
        "humidity, pressure) were computed strictly over trailing history (shifted by 1 day) so current-day values never leaked into historical windows.\n"
        "5. Preprocessing Pipeline: A scikit-learn ColumnTransformer applied StandardScaler to all 32 numerical features and "
        "OneHotEncoder(drop='first', handle_unknown='ignore') to the 3 categorical features (location, weather_code, season), producing 51 transformed inputs."
    )

    # 6. EXPLORATORY DATA ANALYSIS (EDA)
    add_section_header("6", "Exploratory Data Analysis (EDA)")
    doc.add_paragraph(
        "Exploratory data analysis was conducted across all 72,691 records to discover physical weather drivers and patterns:"
    )

    # Helper to insert figure
    fig_counter = 1
    def insert_figure(img_name, caption, width=5.5):
        nonlocal fig_counter
        img_path = FIG_DIR / img_name
        if img_path.exists():
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(2)
            doc.add_picture(str(img_path), width=Inches(width))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(8)
            run_cap = p_cap.add_run(f"Figure {fig_counter}: {caption}")
            run_cap.font.size = Pt(9.5)
            run_cap.font.italic = True
            run_cap.font.color.rgb = RGBColor(90, 90, 90)
            fig_counter += 1

    insert_figure(
        "eda_01_target_distribution.png",
        "Target variable class distribution (42.22% Rainy vs 57.78% Dry days across Maharashtra).",
        width=4.8
    )
    doc.add_paragraph(
        "Observation: The target variable displays a moderate balance (42.22% positive rain instances), confirming "
        "that the dataset is naturally well-conditioned for binary classification without requiring artificial resampling techniques."
    )

    insert_figure(
        "eda_05_monthly_seasonal_patterns.png",
        "Monthly rainfall distribution and seasonal progression in Maharashtra.",
        width=5.5
    )
    doc.add_paragraph(
        "Observation: Rainfall is overwhelmingly concentrated during the Southwest Monsoon (June–September), peaking in "
        "July (96.8% rain occurrence) and August (96.3%), while winter months (December–February) experience <5% rain occurrence."
    )

    insert_figure(
        "eda_07_humidity_vs_rainfall.png",
        "Relative humidity distributions and atmospheric moisture precursors for rainy vs dry days.",
        width=5.2
    )
    doc.add_paragraph(
        "Observation: Mean relative humidity averages 80.2% on days preceding rain versus 56.3% on dry days (+23.9% delta), "
        "confirming that atmospheric boundary layer moisture is the single strongest physical indicator of upcoming rain."
    )

    insert_figure(
        "eda_03_rainfall_by_location.png",
        "Daily rainfall distributions and regional gradients across the 8 Maharashtra locations.",
        width=5.5
    )
    doc.add_paragraph(
        "Observation: Coastal Konkan stations (Ratnagiri and Mumbai) experience significantly higher daily rainfall volume "
        "and rain probability than inland rain-shadow plateau stations (Solapur and Chhatrapati Sambhajinagar)."
    )

    # 7. MODEL DEVELOPMENT
    add_section_header("7", "Model Development")
    doc.add_paragraph(
        "Three distinct supervised classification algorithms were selected to provide a diverse spectrum of learning inductive biases:\n"
        "1. Logistic Regression (Baseline Linear Model): Serves as an interpretable linear benchmark. Optimized with L2 regularization "
        "and liblinear solver to model log-odds of rain occurrence.\n"
        "2. Random Forest (Bagging Ensemble): Builds an ensemble of de-correlated decision trees (max_depth=20, min_samples_leaf=8, "
        "n_estimators=100) using bootstrap aggregation to capture non-linear feature interactions and resist overfitting.\n"
        "3. Gradient Boosting (Sequential Boosting Ensemble): Sequentially trains shallow decision trees (n_estimators=100, max_depth=4, "
        "learning_rate=0.1, subsample=0.8) to minimize log-loss, effectively focusing on hard-to-classify atmospheric transition days."
    )

    # 8. MODEL EVALUATION
    add_section_header("8", "Model Evaluation & Results")
    doc.add_paragraph(
        "All models were evaluated across both the 2023 Validation set (2,920 records) and the 2024 Final Test set (2,920 records). "
        "Evaluation metrics include Accuracy, Precision, Recall, F1-Score, and ROC-AUC:"
    )

    # Results Table
    res_table = doc.add_table(rows=7, cols=7)
    res_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    res_headers = ["Model", "Stage", "Split", "Accuracy", "Precision", "Recall", "ROC-AUC"]
    for j, h in enumerate(res_headers):
        cell = res_table.rows[0].cells[j]
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(9)
        set_cell_background(cell, "1565C0")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell, 60, 60, 80, 80)

    perf_rows = [
        ("Logistic Regression", "Tuned", "Val (2023)", "0.8736", "0.9067", "0.8008", "0.9448"),
        ("Logistic Regression", "Tuned", "Test (2024)", "0.9045", "0.9391", "0.8652", "0.9680"),
        ("Random Forest", "Tuned", "Val (2023)", "0.8702", "0.9087", "0.7901", "0.9468"),
        ("Random Forest", "Tuned", "Test (2024)", "0.9055", "0.9533", "0.8528", "0.9722"),
        ("Gradient Boosting", "Tuned", "Val (2023)", "0.8740", "0.9124", "0.7954", "0.9499"),
        ("Gradient Boosting", "Tuned", "Test (2024)", "0.9062", "0.9506", "0.8569", "0.9714"),
    ]
    for i, row in enumerate(perf_rows, start=1):
        for j, val in enumerate(row):
            cell = res_table.rows[i].cells[j]
            cell.text = val
            cell.paragraphs[0].runs[0].font.size = Pt(8.5)
            if "Gradient Boosting" in row[0]:
                set_cell_background(cell, "E8F5E9")
                cell.paragraphs[0].runs[0].font.bold = True
            else:
                set_cell_background(cell, "F9FBFD" if i % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, 50, 50, 80, 80)

    insert_figure(
        "phase6_baseline_confusion_matrices.png",
        "Confusion matrices for all models on 2023 Validation and 2024 Test sets.",
        width=5.5
    )
    insert_figure(
        "phase6_baseline_roc_curves.png",
        "ROC Curves comparing discrimination performance across models on Validation and Test sets.",
        width=5.5
    )
    doc.add_paragraph(
        "Observation: All three models demonstrate strong discriminative power (ROC-AUC > 0.94 on Validation, > 0.96 on Test). "
        "Test performance in 2024 was higher than 2023 across all algorithms due to stronger, more continuous monsoon signatures during 2024."
    )

    # 9. ERROR ANALYSIS
    add_section_header("9", "Error Analysis")
    doc.add_paragraph(
        "Detailed error analysis on the 2023 Validation partition (N = 2,920) for the top-performing Gradient Boosting model revealed:\n"
        "• Confusion Breakdown: True Positives = 1,042 (35.7%), True Negatives = 1,510 (51.7%), False Positives = 100 (3.4%), False Negatives = 268 (9.2%).\n"
        "• Seasonal Asymmetry: 61.6% of all missed rain days (165 out of 268 FNs) occurred during the Summer / Pre-Monsoon season (April–June), "
        "where the FN rate rose to 22.4%. These errors stem from rapid, localized convective thunderstorms that develop without multi-day synoptic moisture buildup.\n"
        "• Monsoon Accuracy: During the active Southwest Monsoon (July–September), the model was highly accurate with an FN rate of only 2.4% and FP rate of 3.9%."
    )

    insert_figure(
        "phase6_feature_importance.png",
        "Top 20 feature importances for Gradient Boosting and Random Forest, and standardized coefficients for Logistic Regression.",
        width=5.5
    )
    doc.add_paragraph(
        "Observation: Across tree ensembles and logistic regression, rain_sum_lag_1 (yesterday's rain), rain_sum_3d_total "
        "(accumulated wetness), weather_code (WMO state), and relative_humidity_2m_mean are the most influential predictive features."
    )

    # 10. FINAL MODEL
    add_section_header("10", "Final Model Selection")
    doc.add_paragraph(
        "Based on empirical evidence from Phase 6, Gradient Boosting was selected as the final production model (models/best_model_phase6.joblib):\n"
        "1. Top Validation Discrimination: Highest Validation ROC-AUC (0.9499) and Average Precision (0.9388).\n"
        "2. Generalization Stability: Smallest performance delta between Validation and Test (ΔAUC = 0.0215 vs 0.0254 for RF).\n"
        "3. Probability Forecast Quality: Best probability calibration with the lowest Brier score (Val: 0.0897, Test: 0.0665).\n"
        "4. Balanced Test Accuracy: Top Test F1-Score of 0.9014 and Test Accuracy of 90.62%."
    )

    # 11. STREAMLIT APPLICATION
    add_section_header("11", "Streamlit Web Application")
    doc.add_paragraph(
        "A lightweight web interface was developed in app/app.py using Streamlit. The application implements an automated workflow:\n"
        "• Location Selection: The user selects one of the 8 Maharashtra locations from a dropdown menu.\n"
        "• Automated Live Data Retrieval: The app queries the Open-Meteo Forecast API with past_days=8 and timezone=Asia/Kolkata.\n"
        "• Real-Time Feature Pipeline: Transforms raw API payloads into the exact 35-feature schema required by the model.\n"
        "• Instant Inference: The Gradient Boosting pipeline predicts next-day rain (YES / NO) alongside the predicted probability percentage and atmospheric context metrics."
    )

    insert_figure(
        "streamlit_app_screenshot.png",
        "WeatherCast Streamlit web application running live inference.",
        width=5.2
    )

    # 12. RESULTS & DISCUSSION
    add_section_header("12", "Results & Discussion")
    doc.add_paragraph(
        "The project successfully demonstrates that machine learning pipelines can achieve high operational accuracy "
        "(>90% accuracy, >0.97 ROC-AUC) on next-day rain classification using publicly available ERA5 reanalysis data. "
        "Key takeaways include:\n"
        "• Historical Context Matters: Incorporating multi-day lag and rolling features improved predictive stability compared to relying solely on single-day measurements.\n"
        "• Transition Challenges: Convective pre-monsoon showers represent the main source of false negatives, highlighting the inherent physical limits of 24-hour daily aggregate forecasting.\n"
        "• Usability: Eliminating manual parameter inputs in the Streamlit app ensures practical real-world utility."
    )

    # 13. LIMITATIONS
    add_section_header("13", "Limitations")
    doc.add_paragraph(
        "1. Geographic Coverage: Evaluated on 8 key Maharashtra urban centers; micro-climates in dense forested or ghat regions may differ.\n"
        "2. ERA5 Model Resolution: Reanalysis datasets have ~5–15% inherent precipitation uncertainty compared to ground rain gauges.\n"
        "3. Binary Simplification: Predicts rain occurrence (yes/no) rather than quantitative rainfall volume or storm intensity.\n"
        "4. Non-Stationarity: Long-term climate shifts or extreme El Niño events may alter historical statistical correlations."
    )

    # 14. CONCLUSION
    add_section_header("14", "Conclusion")
    doc.add_paragraph(
        "WeatherCast successfully delivers a complete, leak-free, and empirically validated machine learning system for next-day rain prediction "
        "in Maharashtra. Through rigorous chronological validation, feature engineering, and ensemble modeling, Gradient Boosting emerged as the "
        "optimal model, delivering 90.62% test accuracy and 0.9714 test ROC-AUC. The integrated Streamlit application proves the feasibility "
        "of seamless real-time weather intelligence without requiring complex user inputs."
    )

    # 15. REFERENCES
    add_section_header("15", "References")
    doc.add_paragraph(
        "1. Open-Meteo Historical Weather API & Forecast API Documentation: https://open-meteo.com/\n"
        "2. Hersbach, H., et al. (2020). The ERA5 global reanalysis. Quarterly Journal of the Royal Meteorological Society, 146(730), 1999-2049.\n"
        "3. India Meteorological Department (IMD). Standard Monsoon & Seasonal Definitions. Ministry of Earth Sciences, Govt. of India.\n"
        "4. Pedregosa, F., et al. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825-2830.\n"
        "5. Streamlit Documentation: https://docs.streamlit.io/"
    )

    docx_path = DOCS_DIR / "WeatherCast_Final_Report.docx"
    doc.save(str(docx_path))
    print(f"✅ Word Document successfully generated: {docx_path}")
    return docx_path

# ─────────────────────────────────────────────────────────────────────────────
# 2. HTML GENERATION (Print-Ready A4 with CSS & Embedded Images)
# ─────────────────────────────────────────────────────────────────────────────

def get_base64_img(img_name):
    p = FIG_DIR / img_name
    if p.exists():
        with open(p, "rb") as f:
            b64 = base64.b64encode(f.read()).decode("utf-8")
            return f"data:image/png;base64,{b64}"
    return ""

def create_html_report():
    img_target = get_base64_img("eda_01_target_distribution.png")
    img_season = get_base64_img("eda_05_monthly_seasonal_patterns.png")
    img_humidity = get_base64_img("eda_07_humidity_vs_rainfall.png")
    img_loc = get_base64_img("eda_03_rainfall_by_location.png")
    img_cm = get_base64_img("phase6_baseline_confusion_matrices.png")
    img_roc = get_base64_img("phase6_baseline_roc_curves.png")
    img_feat = get_base64_img("phase6_feature_importance.png")
    img_app = get_base64_img("streamlit_app_screenshot.png")

    html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>WeatherCast — Final Project Report (ML Case Study 75)</title>
<style>
  @page {{
    size: A4;
    margin: 20mm 15mm 20mm 15mm;
  }}
  @media print {{
    body {{
      font-size: 10.5pt;
      line-height: 1.4;
      color: #222;
      background: #fff;
    }}
    .no-print {{
      display: none !important;
    }}
    .page-break {{
      page-break-before: always;
    }}
    .figure-container {{
      page-break-inside: avoid;
    }}
    table {{
      page-break-inside: avoid;
    }}
  }}
  body {{
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
    color: #333;
    line-height: 1.5;
    background-color: #f4f6f9;
    margin: 0;
    padding: 0;
  }}
  .report-wrapper {{
    max-width: 860px;
    margin: 20px auto;
    background: #fff;
    padding: 40px 50px;
    box-shadow: 0 4px 15px rgba(0,0,0,0.08);
    border-radius: 4px;
  }}
  .btn-print {{
    position: fixed;
    top: 20px;
    right: 20px;
    background: #1565c0;
    color: #fff;
    border: none;
    padding: 12px 20px;
    font-size: 14px;
    font-weight: bold;
    border-radius: 6px;
    cursor: pointer;
    box-shadow: 0 2px 8px rgba(0,0,0,0.2);
    z-index: 1000;
  }}
  .btn-print:hover {{
    background: #0d47a1;
  }}
  .header-box {{
    text-align: center;
    border-bottom: 2px solid #1565c0;
    padding-bottom: 20px;
    margin-bottom: 25px;
  }}
  h1 {{
    color: #1565c0;
    font-size: 26pt;
    margin: 0 0 8px 0;
  }}
  .subtitle {{
    font-size: 13pt;
    font-weight: 600;
    color: #555;
    margin-bottom: 15px;
  }}
  .meta-table {{
    width: 100%;
    border-collapse: collapse;
    margin: 15px 0 25px 0;
  }}
  .meta-table td {{
    padding: 8px 12px;
    font-size: 10pt;
    border: 1px solid #e0e0e0;
  }}
  .meta-table td:first-child {{
    background: #f0f4f8;
    font-weight: bold;
    width: 35%;
    color: #1565c0;
  }}
  h2 {{
    color: #1565c0;
    font-size: 14pt;
    border-bottom: 1.5px solid #e3edf7;
    padding-bottom: 4px;
    margin-top: 28px;
    margin-bottom: 12px;
  }}
  p, li {{
    font-size: 10.5pt;
    text-align: justify;
    margin-bottom: 10px;
  }}
  ul {{
    margin-top: 5px;
    padding-left: 20px;
  }}
  .data-table {{
    width: 100%;
    border-collapse: collapse;
    margin: 15px 0;
    font-size: 9.5pt;
  }}
  .data-table th {{
    background: #1565c0;
    color: #fff;
    padding: 8px 10px;
    text-align: left;
    font-weight: 600;
  }}
  .data-table td {{
    padding: 7px 10px;
    border-bottom: 1px solid #e0e0e0;
  }}
  .data-table tr:nth-child(even) {{
    background: #f9fbfd;
  }}
  .highlight-row {{
    background: #e8f5e9 !important;
    font-weight: bold;
  }}
  .figure-container {{
    text-align: center;
    margin: 20px 0;
  }}
  .figure-container img {{
    max-width: 90%;
    height: auto;
    border: 1px solid #e0e0e0;
    border-radius: 4px;
    box-shadow: 0 2px 6px rgba(0,0,0,0.05);
  }}
  .caption {{
    font-size: 9pt;
    font-style: italic;
    color: #666;
    margin-top: 6px;
  }}
  .obs-box {{
    background: #f8faff;
    border-left: 4px solid #1565c0;
    padding: 8px 12px;
    margin: 8px 0 16px 0;
    font-size: 9.5pt;
  }}
</style>
</head>
<body>

<button class="btn-print no-print" onclick="window.print()">🖨️ Print / Download as PDF</button>

<div class="report-wrapper">

  <div class="header-box">
    <h1>WeatherCast 🌦️</h1>
    <div class="subtitle">Maharashtra Weather Pattern Analysis & Next-Day Rain Prediction</div>
    <div>B.Tech Semester V — Machine Learning Case Study 75 Final Report</div>
  </div>

  <table class="meta-table">
    <tr>
      <td>Academic Case Study</td>
      <td>Case Study 75 — Weather Pattern Analysis</td>
    </tr>
    <tr>
      <td>ML Problem Formulation</td>
      <td>Supervised Binary Classification (<code>rain_tomorrow</code> ∈ {{0, 1}})</td>
    </tr>
    <tr>
      <td>Dataset Source & Model</td>
      <td>Open-Meteo Historical Weather API (ECMWF ERA5 Reanalysis, 2000–2024)</td>
    </tr>
    <tr>
      <td>Geographic Scope</td>
      <td>8 Maharashtra Cities across Konkan, Madhya Maharashtra, Marathwada & Vidarbha</td>
    </tr>
  </table>

  <h2>1. Abstract</h2>
  <p>
    Weather forecasting plays an indispensable role in agricultural scheduling, flood management, and daily logistics 
    across Maharashtra. This case study develops <strong>WeatherCast</strong>, an end-to-end machine learning system for 
    next-day rainfall prediction across 8 representative Maharashtra cities over a 25-year period (2000–2024, 72,691 observations). 
    A zero-leakage feature-engineering pipeline transformed raw atmospheric variables into 35 base inputs, including cyclical 
    temporal encodings, multi-day lag metrics, and strictly trailing rolling statistics. Three algorithms were tuned using 
    chronological splits (Train: 2000–2022, Val: 2023, Test: 2024): Logistic Regression, Random Forest, and Gradient Boosting. 
    Gradient Boosting achieved the highest overall performance with a Validation ROC-AUC of <strong>0.9499</strong>, 
    Test ROC-AUC of <strong>0.9714</strong>, Test F1-score of <strong>0.9014</strong>, and superior probability calibration (Brier score: 0.0665). 
    An interactive Streamlit application was built that automatically fetches live Open-Meteo atmospheric observations for any 
    selected city, producing real-time rain predictions and calibrated confidence scores without requiring manual user inputs.
  </p>

  <h2>2. Introduction</h2>
  <p>
    Precipitation forecasting in Maharashtra is complicated by sharp topographic gradients and the intense seasonality of the 
    Southwest Monsoon. Annual rainfall ranges from over 3,000 mm along the Western Ghats and coastal Konkan belt to less than 600 mm 
    in the rain-shadow plateau. While classical Numerical Weather Prediction (NWP) models require massive supercomputing clusters, 
    machine learning provides a fast, data-driven approach by capturing statistical relationships between surface atmospheric precursors 
    (such as humidity jumps, barometric pressure drops, and wind shifts) and subsequent precipitation.
  </p>

  <h2>3. Problem Definition</h2>
  <p>
    The project is formulated as a <strong>Supervised Binary Classification</strong> problem:
  </p>
  <ul>
    <li><strong>Input Feature Vector (X):</strong> 35 atmospheric, geographic, cyclical temporal, and lag/rolling features for date <em>t</em>.</li>
    <li><strong>Target Variable (y):</strong> <code>rain_tomorrow</code> ∈ {{0, 1}} for date <em>t+1</em>:
      <ul>
        <li><code>1 (Rain)</code>: Measurable next-day liquid precipitation (<code>rain_sum > 0.0 mm</code>).</li>
        <li><code>0 (No Rain)</code>: Dry next day (<code>rain_sum == 0.0 mm</code>).</li>
      </ul>
    </li>
  </ul>

  <h2>4. Dataset & Data Source</h2>
  <p>
    Data was acquired from the Open-Meteo Historical Weather API based on the peer-reviewed ECMWF ERA5 reanalysis model (0.25° spatial grid). 
    The raw dataset covers 25 continuous calendar years (2000-01-01 to 2024-12-31) across 8 cities:
  </p>

  <table class="data-table">
    <thead>
      <tr>
        <th>Location</th>
        <th>Meteorological Region</th>
        <th>Coordinates</th>
        <th>Rainy Days % (2000–2024)</th>
      </tr>
    </thead>
    <tbody>
      <tr><td>Mumbai</td><td>Konkan (Coastal)</td><td>19.07°N, 72.88°E</td><td>43.4%</td></tr>
      <tr><td>Ratnagiri</td><td>Konkan (Coastal)</td><td>16.99°N, 73.31°E</td><td>46.8%</td></tr>
      <tr><td>Pune</td><td>Madhya Maharashtra</td><td>18.52°N, 73.86°E</td><td>44.9%</td></tr>
      <tr><td>Nashik</td><td>Madhya Maharashtra</td><td>20.00°N, 73.79°E</td><td>42.3%</td></tr>
      <tr><td>Kolhapur</td><td>Madhya Maharashtra</td><td>16.70°N, 74.23°E</td><td>45.1%</td></tr>
      <tr><td>Solapur</td><td>Madhya Maharashtra</td><td>17.67°N, 75.91°E</td><td>35.8%</td></tr>
      <tr><td>Chhatrapati Sambhajinagar</td><td>Marathwada</td><td>19.88°N, 75.34°E</td><td>37.5%</td></tr>
      <tr><td>Nagpur</td><td>Vidarbha</td><td>21.15°N, 79.08°E</td><td>41.9%</td></tr>
    </tbody>
  </table>
  <p>
    <strong>Date-Continuity & Gap Handling:</strong> A strict validation rule ensured <code>rain_tomorrow</code> was created only when 
    <code>next_date == current_date + 1 day</code>. In Chhatrapati Sambhajinagar, the raw archive contained a 1-year gap in 2002 
    (from 2001-12-31 to 2003-01-01); the target on 2001-12-31 was explicitly assigned <code>NaN</code> to prevent bridging across 366 days.
  </p>

  <h2>5. Data Preprocessing & Feature Engineering</h2>
  <p>
    The raw data was transformed into an ML-ready feature matrix following strict zero-leakage protocols:
  </p>
  <ul>
    <li><strong>Chronological Splits:</strong> Train: 2000–2022 (66,779 rows), Validation: 2023 (2,920 rows), Test: 2024 (2,920 rows).</li>
    <li><strong>Cyclical Date Encodings:</strong> Sine/Cosine projections of month and day-of-year (<code>month_sin, month_cos, day_of_year_sin, day_of_year_cos</code>).</li>
    <li><strong>Calendar-Strict Lags:</strong> 8 lag features (t-1, t-3, t-7) computed on full daily calendar grids.</li>
    <li><strong>Backward Trailing Rolling Windows:</strong> 6 aggregations (3-day and 7-day trailing total rain, mean max/min temp, humidity, pressure) shifted by 1 day.</li>
    <li><strong>Preprocessing Pipeline:</strong> Standard scaling on 32 numerical features and One-Hot Encoding on 3 categorical features (<code>location, weather_code, season</code>).</li>
  </ul>

  <h2>6. Exploratory Data Analysis (EDA)</h2>
  <p>Exploratory analysis yielded crucial insights into Maharashtra's weather dynamics:</p>

  <div class="figure-container">
    <img src="{img_target}" alt="Target Distribution">
    <div class="caption">Figure 1: Target variable distribution across the 25-year historical dataset.</div>
  </div>
  <div class="obs-box">
    <strong>Observation:</strong> With 42.22% rainy days (30,687 observations) and 57.78% dry days, the dataset exhibits natural class balance, eliminating the need for synthetic resampling.
  </div>

  <div class="figure-container">
    <img src="{img_season}" alt="Seasonal Patterns">
    <div class="caption">Figure 2: Monthly rainfall volume and occurrence rates across seasons.</div>
  </div>
  <div class="obs-box">
    <strong>Observation:</strong> Rain occurrence is concentrated in the Southwest Monsoon (June–September), reaching peak frequencies in July (96.8%) and August (96.3%), while winter months experience <5% rain occurrence.
  </div>

  <div class="figure-container">
    <img src="{img_humidity}" alt="Humidity vs Rainfall">
    <div class="caption">Figure 3: Relative humidity distributions for rainy versus dry next-day events.</div>
  </div>
  <div class="obs-box">
    <strong>Observation:</strong> Mean relative humidity averages 80.2% prior to rain days versus 56.3% on dry days (+23.9% delta), making boundary-layer moisture the single strongest physical rain precursor.
  </div>

  <div class="figure-container">
    <img src="{img_loc}" alt="Location Patterns">
    <div class="caption">Figure 4: Regional precipitation volume differences across Maharashtra.</div>
  </div>
  <div class="obs-box">
    <strong>Observation:</strong> Coastal Konkan cities (Ratnagiri, Mumbai) receive over 2.4x the daily rainfall volume of inland plateau stations (Solapur, Sambhajinagar).
  </div>

  <h2>7. Model Development</h2>
  <p>Three models were trained and tuned using time-aware validation (<code>TimeSeriesSplit(n_splits=3)</code> on 2000–2022):</p>
  <ul>
    <li><strong>Logistic Regression (Baseline):</strong> Linear model with L2 regularization and liblinear solver (C=1.0).</li>
    <li><strong>Random Forest (Bagging):</strong> 100 de-correlated trees with max_depth=20 and min_samples_leaf=8.</li>
    <li><strong>Gradient Boosting (Sequential Boosting):</strong> 100 sequential trees with max_depth=4, learning_rate=0.1, and subsample=0.8.</li>
  </ul>

  <h2>8. Model Evaluation & Results</h2>
  <p>Performance results across the 2023 Validation set and 2024 Unseen Final Test set:</p>

  <table class="data-table">
    <thead>
      <tr>
        <th>Model</th>
        <th>Stage</th>
        <th>Split</th>
        <th>Accuracy</th>
        <th>Precision</th>
        <th>Recall</th>
        <th>F1-Score</th>
        <th>ROC-AUC</th>
      </tr>
    </thead>
    <tbody>
      <tr><td>Logistic Regression</td><td>Tuned</td><td>Val (2023)</td><td>0.8736</td><td>0.9067</td><td>0.8008</td><td>0.8504</td><td>0.9448</td></tr>
      <tr><td>Logistic Regression</td><td>Tuned</td><td>Test (2024)</td><td>0.9045</td><td>0.9391</td><td>0.8652</td><td>0.9006</td><td>0.9680</td></tr>
      <tr><td>Random Forest</td><td>Tuned</td><td>Val (2023)</td><td>0.8702</td><td>0.9087</td><td>0.7901</td><td>0.8452</td><td>0.9468</td></tr>
      <tr><td>Random Forest</td><td>Tuned</td><td>Test (2024)</td><td>0.9055</td><td>0.9533</td><td>0.8528</td><td>0.9003</td><td>0.9722</td></tr>
      <tr class="highlight-row"><td>Gradient Boosting</td><td>Tuned (Selected)</td><td>Val (2023)</td><td>0.8740</td><td>0.9124</td><td>0.7954</td><td>0.8499</td><td>0.9499</td></tr>
      <tr class="highlight-row"><td>Gradient Boosting</td><td>Tuned (Selected)</td><td>Test (2024)</td><td>0.9062</td><td>0.9506</td><td>0.8569</td><td>0.9014</td><td>0.9714</td></tr>
    </tbody>
  </table>

  <div class="figure-container">
    <img src="{img_cm}" alt="Confusion Matrices">
    <div class="caption">Figure 5: Confusion matrices for all models on 2023 Validation and 2024 Test sets.</div>
  </div>

  <div class="figure-container">
    <img src="{img_roc}" alt="ROC Curves">
    <div class="caption">Figure 6: ROC Curves comparing discrimination performance across models.</div>
  </div>

  <h2>9. Error Analysis</h2>
  <p>
    Error breakdown for Gradient Boosting on the 2023 Validation set (N = 2,920):
  </p>
  <ul>
    <li><strong>False Positives (3.4%, 100 cases):</strong> False alarms occurred during high-humidity borderline days where rain did not materialize.</li>
    <li><strong>False Negatives (9.2%, 268 cases):</strong> Missed rain days were heavily concentrated in the <strong>Summer / Pre-Monsoon</strong> season (165 of 268 cases, 22.4% FN rate), where isolated convective showers develop rapidly without sustained multi-day synoptic moisture buildup.</li>
    <li><strong>Monsoon Stability:</strong> In the active Southwest Monsoon, the FN rate was only 2.4% and FP rate was 3.9%.</li>
  </ul>

  <div class="figure-container">
    <img src="{img_feat}" alt="Feature Importance">
    <div class="caption">Figure 7: Top feature importances and standardized coefficients across models.</div>
  </div>
  <div class="obs-box">
    <strong>Observation:</strong> Yesterday's rainfall (<code>rain_sum_lag_1</code>), 3-day accumulated rainfall (<code>rain_sum_3d_total</code>), WMO weather code, and mean relative humidity are the most influential predictive drivers.
  </div>

  <h2>10. Final Model Selection</h2>
  <p>
    <strong>Gradient Boosting</strong> (<code>models/best_model_phase6.joblib</code>) was selected as the final production model based on:
  </p>
  <ul>
    <li>Highest Validation ROC-AUC (<strong>0.9499</strong>) and Average Precision (<strong>0.9388</strong>).</li>
    <li>Superior generalisation stability (ΔAUC = 0.0215 from validation to test).</li>
    <li>Best probability forecast quality (Validation Brier score: <strong>0.0897</strong>, Test Brier score: <strong>0.0665</strong>).</li>
    <li>Top balanced Test F1-score (<strong>0.9014</strong>) and Test Accuracy (<strong>90.62%</strong>).</li>
  </ul>

  <h2>11. Streamlit Application</h2>
  <p>
    A standalone web application was implemented in <code>app/app.py</code>:
  </p>
  <ul>
    <li><strong>Zero Manual Input:</strong> Users simply choose one of the 8 cities from a dropdown.</li>
    <li><strong>Automated Data Fetching:</strong> The app queries the Open-Meteo Forecast API with <code>past_days=8</code> and <code>timezone=Asia/Kolkata</code>.</li>
    <li><strong>Instant Feature Generation:</strong> Automatically builds all 35 lag, rolling, and cyclical features.</li>
    <li><strong>Real-Time Forecast:</strong> Delivers instant binary prediction (YES / NO), calibrated probability percentage, and atmospheric metrics.</li>
  </ul>

  <div class="figure-container">
    <img src="{img_app}" alt="Streamlit App">
    <div class="caption">Figure 8: WeatherCast Streamlit Web Application Interface.</div>
  </div>

  <h2>12. Results & Discussion</h2>
  <p>
    The results prove that ML models trained on high-quality reanalysis data can accurately forecast daily precipitation 
    without needing supercomputing resources. Incorporating trailing lag and rolling window features significantly improved 
    generalization stability over single-day observations.
  </p>

  <h2>13. Limitations</h2>
  <ul>
    <li><strong>Geographic Scope:</strong> Covers 8 major urban stations; micro-climatic pockets in high-altitude ghats may differ.</li>
    <li><strong>Reanalysis Resolution:</strong> ERA5 has ~5–15% inherent precipitation uncertainty compared to ground rain gauges.</li>
    <li><strong>Binary Simplification:</strong> Predicts rain occurrence rather than exact millimeters of rainfall volume.</li>
  </ul>

  <h2>14. Conclusion</h2>
  <p>
    WeatherCast delivers an end-to-end, leak-free, and validated machine learning pipeline for Maharashtra next-day rain prediction. 
    Gradient Boosting demonstrated outstanding performance with 90.62% test accuracy and 0.9714 test ROC-AUC. 
    The Streamlit prototype brings this ML intelligence directly to users through an intuitive 1-click interface.
  </p>

  <h2>15. References</h2>
  <p>
    1. Open-Meteo Weather API Documentation: <a href="https://open-meteo.com/">https://open-meteo.com/</a><br>
    2. Hersbach, H., et al. (2020). The ERA5 global reanalysis. <em>Quarterly Journal of the Royal Meteorological Society</em>, 146(730), 1999-2049.<br>
    3. India Meteorological Department (IMD). Standard Monsoon Definitions. Ministry of Earth Sciences, Govt. of India.<br>
    4. Pedregosa, F., et al. (2011). Scikit-learn: Machine Learning in Python. <em>JMLR</em>, 12, 2825-2830.<br>
    5. Streamlit Framework Documentation: <a href="https://docs.streamlit.io/">https://docs.streamlit.io/</a>
  </p>

</div>

</body>
</html>
"""
    html_path = DOCS_DIR / "WeatherCast_Final_Report.html"
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html_content)
    print(f"✅ HTML Report successfully generated: {html_path}")
    return html_path

if __name__ == "__main__":
    docx_file = create_docx_report()
    html_file = create_html_report()
    print("\nAll Final Report artifacts generated successfully!")
