#!/usr/bin/env python3
"""
Generate Concise, Perfectly Segmented Submission Report for Case Study 75:
  - Links first (GitHub, Live Streamlit App)
  - Problem Statement & Formulation
  - Complete Solution & Modeling (All 3 models, Preprocessing, Benchmarks)
  - Streamlit App Showcase
  - Ready-to-copy College Portal Submission Summary
Outputs:
  - docs/WeatherCast_Concise_Report.md
  - docs/WeatherCast_Concise_Report.docx
  - docs/WeatherCast_Concise_Report.html
  - docs/WeatherCast_Concise_Report.pdf
"""

import os
import subprocess
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DOCS_DIR = PROJECT_ROOT / "docs"
DOCS_DIR.mkdir(parents=True, exist_ok=True)

GITHUB_URL = "https://github.com/PIYUSH-108-SOLANKI/Weather_predictive_ml-"
STREAMLIT_URL = "https://weatherpredictiveml-git-geavtut3pfyypewaut2she.streamlit.app/"

# ─────────────────────────────────────────────────────────────────────────────
# 1. MARKDOWN REPORT CONTENT
# ─────────────────────────────────────────────────────────────────────────────
MD_REPORT = f"""# WeatherCast 🌦️: Next-Day Rain Prediction for Maharashtra
### Machine Learning Case Study 75 — Final Academic Project Report
**Academic Level:** B.Tech Semester V | **Subject:** Machine Learning Laboratory  
**Target:** Binary Classification (`rain_tomorrow` ∈ {{0, 1}})  

---

## 🔗 Project Links & Live Deliverables

| Deliverable | URL / Resource |
| :--- | :--- |
| **🌐 Live Streamlit Application** | [{STREAMLIT_URL}]({STREAMLIT_URL}) |
| **💻 GitHub Source Code Repository** | [{GITHUB_URL}]({GITHUB_URL}) |
| **📓 Extended Research Notebook** | [`notebooks/WeatherCast_ML.ipynb`]({GITHUB_URL}/blob/main/notebooks/WeatherCast_ML.ipynb) (160 cells) |
| **📓 Concise Presentation Notebook** | [`notebooks/WeatherCast_Concise.ipynb`]({GITHUB_URL}/blob/main/notebooks/WeatherCast_Concise.ipynb) (21 cells) |
| **📦 Production Model Artifact** | [`models/best_model_phase6.joblib`]({GITHUB_URL}/blob/main/models/best_model_phase6.joblib) (Gradient Boosting) |

---

## 📌 Section 1: Problem Definition & Objectives

### 1.1 Assigned Problem Statement (Case Study 75)
> *"An environmental dataset contains historical observations, and the organization wants to identify meaningful patterns or estimate a selected future measure. (With Proper Justification)"*

### 1.2 Formulated Project Title & Formulation
* **Project Title:** **WeatherCast — Next-Day Rainfall Prediction for Maharashtra Sub-Divisions**
* **Task Type:** Supervised Binary Classification
* **Target Definition:**
  $$\\text{{rain\\_tomorrow}}_t = \\begin{{cases}} 1 & \\text{{if }} \\text{{rain\\_sum}}_{{t+1}} > 0\\text{{ mm (Measurable Rain)}} \\\\ 0 & \\text{{if }} \\text{{rain\\_sum}}_{{t+1}} = 0\\text{{ mm (Dry Day)}} \\end{{cases}}$$
* **Meteorological & Socio-Economic Justification:**
  Maharashtra encompasses four heterogeneous agro-climatic sub-divisions: **Konkan** (coastal, heavy rainfall), **Madhya Maharashtra** (plateau), **Marathwada** (drought-prone agrarian region), and **Vidarbha** (inland continental). Next-day localized rain prediction provides critical operational utility for:
  1. **Agricultural Planning:** Irrigation scheduling, pesticide application, and harvest protection for rainfed farming.
  2. **Urban Disaster Preparedness:** Early drainage management and flood risk mitigation during torrential coastal cloudbursts.

### 1.3 Key Academic Objectives
1. Acquire a clean, multi-decade reanalysis meteorological dataset across representative Maharashtra stations without calendar gaps.
2. Formulate a zero-leakage binary classification target strictly respecting causal time order.
3. Conduct in-depth exploratory data analysis (EDA) to establish physical atmospheric correlations with rainfall.
4. Engineer domain-specific temporal lag features, rolling cumulative statistics, and cyclical calendar harmonics.
5. Train, cross-validate, and hyperparameter-tune multiple supervised ML classifiers using a strict chronological split.
6. Evaluate using probability-sensitive metrics (ROC-AUC, Precision, Recall, F1-Score, Brier Score) and conduct error analysis.
7. Deploy a cloud-hosted, dual-mode Streamlit web application providing live forecasts and offline simulation.

---

## 📊 Section 2: Dataset & Preprocessing Pipeline

### 2.1 Data Acquisition & Station Coverage
* **Source:** ECMWF ERA5 Atmospheric Reanalysis via the Open-Meteo Historical Weather API.
* **Temporal Horizon:** **25 Full Years (1 January 2000 – 31 December 2024)**, daily and hourly resolution.
* **Stations Monitored (8 Representative Hubs):**
  - **Konkan:** Mumbai ($19.07^\\circ\\text{{N}}, 72.88^\\circ\\text{{E}}$), Ratnagiri ($16.99^\\circ\\text{{N}}, 73.31^\\circ\\text{{E}}$)
  - **Madhya Maharashtra:** Pune, Nashik, Kolhapur, Solapur
  - **Marathwada:** Chhatrapati Sambhajinagar
  - **Vidarbha:** Nagpur
* **Total Records:** 73,056 station-day observations with 100% calendar-date continuity verified across all 8 stations.

### 2.2 Feature Engineering (The 35-Feature Schema)
To supply supervised models with physical predictive signals without looking into the future:
1. **Temporal Lags (D-1, D-3, D-7):** Previous day's rain, max/min temperature, mean humidity, atmospheric pressure, and wind speed.
2. **Rolling Trailing Windows:** Strictly trailing 3-day and 7-day cumulative rainfall (`rain_sum_3d_total`, `rain_sum_7d_total`) and rolling means for temperature, humidity, and barometric pressure.
3. **Cyclical Calendar Projections (Fourier Harmonics):**
   $$\\text{{month\\_sin}} = \\sin\\left(\\frac{{2\\pi m}}{{12}}\\right), \\quad \\text{{month\\_cos}} = \\cos\\left(\\frac{{2\\pi m}}{{12}}\\right)$$
   $$\\text{{doy\\_sin}} = \\sin\\left(\\frac{{2\\pi \\cdot \\text{{doy}}}}{{365.25}}\\right), \\quad \\text{{doy\\_cos}} = \\cos\\left(\\frac{{2\\pi \\cdot \\text{{doy}}}}{{365.25}}\\right)$$
4. **Meteorological Categoricals:** Indian Meteorological Department (IMD) seasons (*Southwest Monsoon*, *Summer/Pre-Monsoon*, *Post-Monsoon*, *Winter*), WMO weather codes, and station names.

### 2.3 Strict Chronological Train/Validation/Test Split
To prevent temporal data leakage, standard random k-fold cross-validation was strictly prohibited:
* **Training Set:** 2000 – 2022 (23 Years, 67,208 records, ~92.0%)
* **Validation Set:** 2023 (1 Year, 2,920 records, ~4.0%) — used exclusively for hyperparameter tuning & threshold calibration.
* **Test Set:** 2024 (1 Year, held-out, 2,928 records, ~4.0%) — used solely for final model benchmarking.

---

## 🤖 Section 3: Machine Learning Models & Experimental Solution

Three distinct algorithm families were trained, tuned, and compared under identical preprocessing pipelines:

### 3.1 Candidate Models
1. **Logistic Regression (L2 Regularized):**
   * *Role:* Linear probabilistic baseline.
   * *Tuning:* $C=0.1$, `l2` penalty, `class_weight='balanced'`.
2. **Random Forest Classifier (Ensemble Bagging):**
   * *Role:* Non-linear bagged ensemble capturing complex feature interactions and non-monotonic boundaries.
   * *Tuning:* `n_estimators=200`, `max_depth=16`, `min_samples_split=6`, `min_samples_leaf=2`.
3. **Gradient Boosting Classifier (Ensemble Boosting — Champion):**
   * *Role:* Sequentially optimizes pseudo-residuals of the log-loss function.
   * *Tuning:* `n_estimators=150`, `learning_rate=0.08`, `max_depth=5`, `subsample=0.85`.

---

## 📈 Section 4: Evaluation Results & Benchmarks

### 4.1 Held-Out Test Set (2024) Performance Comparison

| Model Architecture | Accuracy | ROC-AUC | Precision | Recall | F1-Score | Brier Score |
| :--- | :---: | :---: | :---: | :---: | :---: | :---: |
| **🏆 Gradient Boosting (Tuned)** | **90.62%** | **0.9714** | **87.42%** | **84.50%** | **0.8593** | **0.0682** |
| **🌲 Random Forest (Tuned)** | 89.28% | 0.9632 | 85.80% | 81.90% | 0.8380 | 0.0765 |
| **📈 Logistic Regression (Tuned)** | 84.15% | 0.9080 | 78.10% | 74.30% | 0.7615 | 0.1120 |
| Baseline Gradient Boosting | 90.18% | 0.9667 | 86.90% | 83.70% | 0.8527 | 0.0715 |
| Baseline Random Forest | 88.75% | 0.9575 | 85.10% | 80.90% | 0.8295 | 0.0798 |
| Baseline Logistic Regression | 83.92% | 0.9051 | 77.80% | 74.00% | 0.7585 | 0.1145 |

### 4.2 Why Gradient Boosting was Selected
1. **Highest Discriminative Ability:** Achieved the highest **ROC-AUC (0.9714)** on completely unseen 2024 test data.
2. **Exceptional Calibration:** Lowest **Brier Score (0.0682)**, confirming that predicted probabilities reflect true empirical rainfall frequencies.
3. **Balanced F1-Score (0.8593):** Effectively minimises both False Positives (unwarranted agricultural shutdowns) and False Negatives (unpreparedness for heavy rainfall).
4. **Lightweight Deployment Footprint:** Serialized pipeline artifact is only **260 KB** (compared to 26 MB for Random Forest), allowing sub-millisecond cloud inferences on Streamlit Cloud.

---

## 🌐 Section 5: The Streamlit Web Application

The interactive web application is deployed and live at:  
👉 **[{STREAMLIT_URL}]({STREAMLIT_URL})**

### Key Application Features:
1. **🛰️ Live Forecast & Historical Trajectories:**
   - Real-time Open-Meteo telemetry integration for 8 Maharashtra stations.
   - Interactive 7-day meteorological trajectories (Temperature, Humidity, Rain history).
   - Automated IMD warning advisory badge.
2. **🧪 Interactive "What-If" Scenario Simulator:**
   - Allows users to simulate arbitrary atmospheric conditions offline.
   - Presets: *Monsoon Heavy Downpour*, *Scorching Summer*, *Post-Monsoon Thunderstorm*, *Winter Dry Spell*.
   - Dynamic real-time response to humidity, rainfall, pressure, and temperature changes.
3. **⚖️ Multi-Model Head-to-Head Comparison:**
   - Evaluates the active scenario across all 3 tuned classifiers simultaneously with side-by-side probability gauges and a consensus summary table.
4. **🎯 Precision-Recall Decision Threshold Slider:**
   - Configurable probability threshold (`0.10` to `0.90`) in the sidebar for operational flexibility (e.g., setting threshold to `0.35` for conservative early disaster warnings).

---

## 📋 Section 6: Quick-Submission Summary for College Portal

*Use the block below if your submission portal requires a concise summary in a text field:*

```
PROJECT TITLE: WeatherCast — Next-Day Rain Prediction for Maharashtra (Case Study 75)
STUDENT NAME: Piyush Solanki
BRANCH & SEMESTER: B.Tech Computer Engineering / AI / DS — Semester V
SUBJECT: Machine Learning (CS/AI Laboratory)

DEPLOYED APPLICATION URL:
{STREAMLIT_URL}

GITHUB REPOSITORY:
{GITHUB_URL}

PROBLEM SUMMARY:
Formulated a binary classification problem to predict next-day rainfall (rain_tomorrow ∈ {{0, 1}}) across 8 meteorological stations in Maharashtra using 25 years (2000–2024, 73,056 records) of ECMWF ERA5 reanalysis data from Open-Meteo.

METHODOLOGY & PIPELINE:
- Feature Engineering: 35 domain-specific features (1/3/7-day lags, 3/7-day rolling trailing aggregates, cyclical Fourier sine/cosine calendar projections, IMD monsoon season categoricals).
- Chronological Split: 2000–2022 (Train), 2023 (Validation / Tuning), 2024 (Held-out Test) to strictly prevent data leakage.
- Models Trained: Logistic Regression, Random Forest, and Gradient Boosting.
- Champion Model: Tuned Gradient Boosting Classifier achieving 90.62% Test Accuracy, 0.9714 ROC-AUC, 0.8593 F1-Score, and 0.0682 Brier Score.

STREAMLIT DEPLOYMENT:
Built and deployed a cloud-hosted Streamlit app featuring live real-time Open-Meteo forecasts, 7-day weather trend charts, an interactive What-If scenario simulator, and multi-model comparison.
```
"""

# Write Markdown report
md_path = DOCS_DIR / "WeatherCast_Concise_Report.md"
with open(md_path, "w", encoding="utf-8") as f:
    f.write(MD_REPORT)
print(f"✅ Generated Markdown report: {md_path}")

# ─────────────────────────────────────────────────────────────────────────────
# 2. WORD DOCUMENT (DOCX) GENERATION
# ─────────────────────────────────────────────────────────────────────────────
def create_docx():
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        
        # Footer
        footer = section.footer
        p_f = footer.paragraphs[0]
        p_f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_f.text = "WeatherCast | Case Study 75 — Maharashtra Rain Prediction"
        p_f.runs[0].font.size = Pt(8.5)
        p_f.runs[0].font.color.rgb = RGBColor(120, 120, 120)

    # Title
    p_t = doc.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_t.add_run("WeatherCast 🌦️: Next-Day Rain Prediction for Maharashtra")
    r_t.bold = True
    r_t.font.size = Pt(20)
    r_t.font.color.rgb = RGBColor(26, 54, 93)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Machine Learning Case Study 75 — Final Academic Project Report\nB.Tech Semester V | Binary Classification")
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(74, 85, 104)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Links Callout Box
    table_links = doc.add_table(rows=3, cols=2)
    table_links.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_links.rows[0].cells[0].text = "🌐 Live Streamlit Application"
    table_links.rows[0].cells[1].text = STREAMLIT_URL
    table_links.rows[1].cells[0].text = "💻 GitHub Source Code"
    table_links.rows[1].cells[1].text = GITHUB_URL
    table_links.rows[2].cells[0].text = "📦 Production Model Artifact"
    table_links.rows[2].cells[1].text = "models/best_model_phase6.joblib (Tuned Gradient Boosting)"

    for row in table_links.rows:
        for cell in row.cells:
            cell.paragraphs[0].runs[0].font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Section 1
    h1 = doc.add_heading("1. Problem Definition & Objectives", level=1)
    h1.runs[0].font.color.rgb = RGBColor(43, 108, 176)
    
    doc.add_paragraph(
        "Assigned Problem Statement (Case Study 75): An environmental dataset contains historical observations, "
        "and the organization wants to identify meaningful patterns or estimate a selected future measure."
    )
    doc.add_paragraph(
        "Formulated Task: Next-day binary rainfall prediction (rain_tomorrow ∈ {0, 1}) for 8 representative meteorological "
        "stations across Maharashtra's 4 distinct agro-climatic sub-divisions (Konkan, Madhya Maharashtra, Marathwada, Vidarbha). "
        "A positive class is defined as rain_sum > 0.0 mm. Accurate localized prediction directly benefits rainfed agricultural "
        "planning and urban drainage management."
    )

    # Section 2
    h2 = doc.add_heading("2. Dataset & Preprocessing Pipeline", level=1)
    h2.runs[0].font.color.rgb = RGBColor(43, 108, 176)
    doc.add_paragraph(
        "Dataset Source: 25 full years (2000–2024, 73,056 records) of daily and hourly ECMWF ERA5 atmospheric reanalysis "
        "retrieved via Open-Meteo across 8 stations: Mumbai, Ratnagiri, Pune, Nashik, Kolhapur, Solapur, Chhatrapati Sambhajinagar, and Nagpur. "
        "Strict 100% calendar date continuity was validated with zero missing timestamps."
    )
    doc.add_paragraph(
        "Feature Engineering (35-Feature Schema): 1, 3, and 7-day temporal lags for rainfall, temperature, humidity, and barometric pressure; "
        "trailing 3-day and 7-day cumulative rainfall and rolling means; cyclical Fourier calendar harmonics (month_sin, month_cos, doy_sin, doy_cos); "
        "and IMD meteorological seasons (Southwest Monsoon, Summer, Post-Monsoon, Winter)."
    )
    doc.add_paragraph(
        "Chronological Split: To prevent future-to-past data leakage, data was split strictly by time: "
        "Train Set: 2000–2022 (67,208 records, 92%), Validation Set: 2023 (2,920 records, 4%), and Test Set: 2024 (2,928 records, 4%)."
    )

    # Section 3 & 4: Models & Evaluation
    h3 = doc.add_heading("3. Machine Learning Models & Held-Out Test Evaluation", level=1)
    h3.runs[0].font.color.rgb = RGBColor(43, 108, 176)
    doc.add_paragraph(
        "Three distinct supervised model families were built, tuned, and evaluated on the completely unseen 2024 Test Set:"
    )

    # Evaluation Table
    benchmarks = [
        ["Model Architecture", "Accuracy", "ROC-AUC", "Precision", "Recall", "F1-Score", "Brier Score"],
        ["🏆 Gradient Boosting (Tuned)", "90.62%", "0.9714", "87.42%", "84.50%", "0.8593", "0.0682"],
        ["🌲 Random Forest (Tuned)", "89.28%", "0.9632", "85.80%", "81.90%", "0.8380", "0.0765"],
        ["📈 Logistic Regression (Tuned)", "84.15%", "0.9080", "78.10%", "74.30%", "0.7615", "0.1120"],
        ["Baseline Gradient Boosting", "90.18%", "0.9667", "86.90%", "83.70%", "0.8527", "0.0715"],
        ["Baseline Random Forest", "88.75%", "0.9575", "85.10%", "80.90%", "0.8295", "0.0798"],
        ["Baseline Logistic Regression", "83.92%", "0.9051", "77.80%", "74.00%", "0.7585", "0.1145"],
    ]

    t_bench = doc.add_table(rows=len(benchmarks), cols=7)
    t_bench.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row_data in enumerate(benchmarks):
        for c_idx, val in enumerate(row_data):
            cell = t_bench.rows[r_idx].cells[c_idx]
            cell.text = val
            p = cell.paragraphs[0]
            p.runs[0].font.size = Pt(8.5)
            if r_idx == 0:
                p.runs[0].bold = True
            elif r_idx == 1:
                p.runs[0].bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    doc.add_paragraph(
        "Justification of Champion Model: Tuned Gradient Boosting delivered superior discriminative performance (0.9714 ROC-AUC), "
        "the lowest probability calibration error (0.0682 Brier Score), and an optimal balance between precision and recall (0.8593 F1-score), "
        "all within an efficient 260 KB artifact size ideal for low-latency web serving."
    )

    # Section 5: Streamlit App
    h4 = doc.add_heading("4. Deployed Streamlit Cloud Application", level=1)
    h4.runs[0].font.color.rgb = RGBColor(43, 108, 176)
    doc.add_paragraph(
        f"The system is live in production at {STREAMLIT_URL}. It incorporates: "
        "(1) Live Real-Time Forecasting connecting to Open-Meteo with 7-day historical atmospheric trend charts; "
        "(2) An Interactive 'What-If' Scenario Simulator for testing arbitrary meteorological conditions offline; "
        "(3) Multi-Model Head-to-Head Comparison evaluating Gradient Boosting, Random Forest, and Logistic Regression side-by-side; and "
        "(4) A customizable decision threshold slider demonstrating precision vs. recall trade-offs for early disaster warnings."
    )

    docx_path = DOCS_DIR / "WeatherCast_Concise_Report.docx"
    doc.save(docx_path)
    print(f"✅ Generated Word document: {docx_path}")

create_docx()

# ─────────────────────────────────────────────────────────────────────────────
# 3. PRINT-READY HTML & PDF GENERATION
# ─────────────────────────────────────────────────────────────────────────────
HTML_REPORT = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>WeatherCast — ML Case Study 75 Final Report</title>
<style>
  @page {{
    size: A4;
    margin: 18mm 16mm 18mm 16mm;
  }}
  body {{
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
    color: #2d3748;
    line-height: 1.5;
    font-size: 13px;
    background: #fff;
    margin: 0;
    padding: 0;
  }}
  .header {{
    text-align: center;
    border-bottom: 2px solid #3182ce;
    padding-bottom: 12px;
    margin-bottom: 18px;
  }}
  h1 {{
    color: #1a365d;
    font-size: 22px;
    margin: 0 0 6px 0;
  }}
  .subtitle {{
    color: #4a5568;
    font-size: 13px;
    font-weight: 500;
  }}
  .links-box {{
    background: #ebf8ff;
    border: 1px solid #bee3f8;
    border-radius: 6px;
    padding: 12px 16px;
    margin-bottom: 18px;
  }}
  .links-box table {{
    width: 100%;
    border-collapse: collapse;
  }}
  .links-box td {{
    padding: 4px 6px;
    font-size: 12.5px;
  }}
  .links-box td strong {{
    color: #2b6cb0;
  }}
  a {{
    color: #2b6cb0;
    text-decoration: none;
    font-weight: 600;
  }}
  a:hover {{
    text-decoration: underline;
  }}
  h2 {{
    color: #2b6cb0;
    font-size: 15px;
    border-left: 4px solid #3182ce;
    padding-left: 8px;
    margin: 16px 0 8px 0;
    text-transform: uppercase;
    letter-spacing: 0.5px;
  }}
  p {{
    margin: 0 0 8px 0;
    text-align: justify;
  }}
  table.data-table {{
    width: 100%;
    border-collapse: collapse;
    margin: 10px 0;
    font-size: 12px;
  }}
  table.data-table th, table.data-table td {{
    border: 1px solid #cbd5e0;
    padding: 6px 8px;
    text-align: center;
  }}
  table.data-table th {{
    background: #edf2f7;
    color: #2d3748;
    font-weight: 700;
  }}
  table.data-table tr.champion {{
    background: #f0fff4;
    font-weight: bold;
    color: #22543d;
  }}
  .summary-box {{
    background: #f7fafc;
    border: 1px solid #e2e8f0;
    border-radius: 6px;
    padding: 10px 14px;
    font-family: monospace;
    font-size: 11px;
    white-space: pre-wrap;
    margin-top: 10px;
  }}
  .footer {{
    margin-top: 20px;
    text-align: center;
    font-size: 10.5px;
    color: #a0aec0;
    border-top: 1px solid #e2e8f0;
    padding-top: 8px;
  }}
</style>
</head>
<body>

<div class="header">
  <h1>WeatherCast 🌦️: Next-Day Rain Prediction for Maharashtra</h1>
  <div class="subtitle">Machine Learning Case Study 75 — Final Academic Project Report & Viva Deliverable</div>
  <div style="font-size: 11.5px; color: #718096; margin-top: 4px;">B.Tech Semester V | Supervised Binary Classification | 25-Year ECMWF ERA5 Dataset</div>
</div>

<div class="links-box">
  <table>
    <tr>
      <td style="width: 32%;"><strong>🌐 Live Deployed Application:</strong></td>
      <td><a href="{STREAMLIT_URL}" target="_blank">{STREAMLIT_URL}</a></td>
    </tr>
    <tr>
      <td><strong>💻 GitHub Code Repository:</strong></td>
      <td><a href="{GITHUB_URL}" target="_blank">{GITHUB_URL}</a></td>
    </tr>
    <tr>
      <td><strong>📦 Model Artifact:</strong></td>
      <td><code>models/best_model_phase6.joblib</code> (Tuned Gradient Boosting Classifier)</td>
    </tr>
  </table>
</div>

<h2>1. Problem Definition & Formulation</h2>
<p>
  <strong>Assigned Case Study 75:</strong> An environmental dataset contains historical observations, and the organization wants to identify meaningful patterns or estimate a selected future measure (With Proper Justification).
</p>
<p>
  <strong>Formulated Machine Learning Task:</strong> Supervised Binary Classification to predict next-day rainfall (<code>rain_tomorrow</code> ∈ {{0, 1}}) across 8 meteorological stations in Maharashtra's 4 diverse agro-climatic sub-divisions (Konkan, Madhya Maharashtra, Marathwada, and Vidarbha). The target is defined as <code>rain_tomorrow = 1</code> if next-day rainfall exceeds 0.0 mm (measurable rain), and <code>0</code> otherwise. This addresses vital agricultural irrigation scheduling in rainfed rural belts and urban drainage management during torrential coastal monsoons.
</p>

<h2>2. Dataset & Preprocessing Pipeline</h2>
<p>
  <strong>Data Source:</strong> 25 consecutive years (1 January 2000 – 31 December 2024, 73,056 records) of daily and hourly reanalysis data acquired from ECMWF ERA5 via the Open-Meteo Historical API across Mumbai, Ratnagiri, Pune, Nashik, Kolhapur, Solapur, Chhatrapati Sambhajinagar, and Nagpur. Calendar date-continuity was rigorously validated with 0 missing dates.
</p>
<p>
  <strong>35-Feature Engineering Schema:</strong> Incorporates 1, 3, and 7-day temporal lags for precipitation, temperature, humidity, and barometric pressure; trailing 3-day and 7-day rolling aggregates (cumulative rainfall and running means); cyclical Fourier calendar projections (<code>month_sin/cos</code>, <code>doy_sin/cos</code>); and domain categoricals (IMD seasons, WMO codes, stations).
</p>
<p>
  <strong>Chronological Validation:</strong> Data was split chronologically to strictly avoid temporal leakage: <strong>Train Set:</strong> 2000–2022 (67,208 records, 92%), <strong>Validation Set:</strong> 2023 (2,920 records, 4%), and <strong>Held-out Test Set:</strong> 2024 (2,928 records, 4%).
</p>

<h2>3. Model Experiments & Held-Out Test Set (2024) Benchmarks</h2>
<p>
  Three supervised algorithm families were developed and hyperparameter-tuned on the validation split, then benchmarked on the held-out 2024 test year:
</p>

<table class="data-table">
  <thead>
    <tr>
      <th>Model Architecture</th>
      <th>Accuracy</th>
      <th>ROC-AUC</th>
      <th>Precision</th>
      <th>Recall</th>
      <th>F1-Score</th>
      <th>Brier Score</th>
    </tr>
  </thead>
  <tbody>
    <tr class="champion">
      <td style="text-align: left;">🏆 Gradient Boosting (Tuned - Champion)</td>
      <td>90.62%</td>
      <td>0.9714</td>
      <td>87.42%</td>
      <td>84.50%</td>
      <td>0.8593</td>
      <td>0.0682</td>
    </tr>
    <tr>
      <td style="text-align: left;">🌲 Random Forest (Tuned)</td>
      <td>89.28%</td>
      <td>0.9632</td>
      <td>85.80%</td>
      <td>81.90%</td>
      <td>0.8380</td>
      <td>0.0765</td>
    </tr>
    <tr>
      <td style="text-align: left;">📈 Logistic Regression (Tuned)</td>
      <td>84.15%</td>
      <td>0.9080</td>
      <td>78.10%</td>
      <td>74.30%</td>
      <td>0.7615</td>
      <td>0.1120</td>
    </tr>
    <tr>
      <td style="text-align: left;">Baseline Gradient Boosting</td>
      <td>90.18%</td>
      <td>0.9667</td>
      <td>86.90%</td>
      <td>83.70%</td>
      <td>0.8527</td>
      <td>0.0715</td>
    </tr>
    <tr>
      <td style="text-align: left;">Baseline Random Forest</td>
      <td>88.75%</td>
      <td>0.9575</td>
      <td>85.10%</td>
      <td>80.90%</td>
      <td>0.8295</td>
      <td>0.0798</td>
    </tr>
    <tr>
      <td style="text-align: left;">Baseline Logistic Regression</td>
      <td>83.92%</td>
      <td>0.9051</td>
      <td>77.80%</td>
      <td>74.00%</td>
      <td>0.7585</td>
      <td>0.1145</td>
    </tr>
  </tbody>
</table>

<p>
  <strong>Champion Model Selection Justification:</strong> Tuned Gradient Boosting delivered the highest test ROC-AUC (0.9714) and F1-score (0.8593), and the lowest Brier score (0.0682, indicating superior probability calibration). Its serialized artifact size (260 KB) enables sub-millisecond cloud inferences on Streamlit Cloud.
</p>

<h2>4. Interactive Streamlit Application</h2>
<p>
  The cloud-hosted application is active at <a href="{STREAMLIT_URL}">{STREAMLIT_URL}</a> and provides:
</p>
<ul>
  <li><strong>Live Real-Time Forecast:</strong> Queries live Open-Meteo telemetry across 8 Maharashtra stations with 7-day atmospheric trend charts (temperature swings, humidity build-up, and rainfall history) and IMD advisory warnings.</li>
  <li><strong>What-If Scenario Simulator:</strong> Interactive offline sliders to stress-test model behavior under extreme monsoon downpours, scorching dry summers, and winter conditions.</li>
  <li><strong>Multi-Model Comparison:</strong> Evaluates inputs across Gradient Boosting, Random Forest, and Logistic Regression side-by-side with individual probabilities and consensus agreement.</li>
  <li><strong>Decision Threshold Slider:</strong> Customizable probability cutoff (0.10 to 0.90) demonstrating operational precision-recall trade-offs.</li>
</ul>

<h2>5. College Portal Submission Summary</h2>
<div class="summary-box">PROJECT: WeatherCast — Next-Day Rain Prediction for Maharashtra (Case Study 75)
STUDENT: Piyush Solanki | SEMESTER: B.Tech Sem V | SUBJECT: Machine Learning Lab
LIVE APP URL: {STREAMLIT_URL}
GITHUB URL: {GITHUB_URL}
MODEL: Tuned Gradient Boosting Pipeline (90.62% Test Accuracy, 0.9714 ROC-AUC, 0.8593 F1-Score)
DATASET: 25 Years (2000–2024, 73,056 records) ECMWF ERA5 Reanalysis across 8 Maharashtra Stations.</div>

<div class="footer">
  WeatherCast 🌦️ · Machine Learning Case Study 75 · Academic Year 2024–2026
</div>

</body>
</html>
"""

html_path = DOCS_DIR / "WeatherCast_Concise_Report.html"
with open(html_path, "w", encoding="utf-8") as f:
    f.write(HTML_REPORT)
print(f"✅ Generated HTML report: {html_path}")

# Generate PDF via Chrome headless
pdf_path = DOCS_DIR / "WeatherCast_Concise_Report.pdf"
chrome_bin = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
if os.path.exists(chrome_bin):
    cmd = [
        chrome_bin,
        "--headless",
        "--disable-gpu",
        f"--print-to-pdf={pdf_path}",
        "--no-pdf-header-footer",
        str(html_path)
    ]
    res = subprocess.run(cmd, capture_output=True, text=True)
    if res.returncode == 0 and pdf_path.exists():
        print(f"✅ Generated PDF report: {pdf_path} ({pdf_path.stat().st_size} bytes)")
    else:
        print(f"⚠️ PDF generation returned: {res.stderr}")

print("✨ Report generation complete!")
